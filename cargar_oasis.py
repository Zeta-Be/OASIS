#!/usr/bin/env python3
"""
OASIS v2.0 — Carga desde Google Sheets hacia Supabase, un mercado o todos.

Vocabulario: una COMPRA es un paquete de OOH con ID propio (001, 002...),
equivale a una fila del Flow. Un MATERIAL es lo que corre sobre esa compra
(001-A, 001-B...), equivale a una fila del Tracking.

Regla general: TODO SE CARGA. Lo incompleto o inconsistente no se deja fuera:
entra, y v_alertas lo muestra en Alertas de su mercado hasta que alguien lo
corrija. Lo único que no entra es lo que la base no puede guardar (un material
cuya compra no existe en el Flow) y las filas del Flow que solo tienen el ID.

Se puede correr tantas veces como haga falta: hace upsert por llave natural.

Uso:
    source .env
    python3 cargar_oasis.py --mercado MLC --dry-run   # simula un mercado
    python3 cargar_oasis.py --todos --dry-run         # simula los de mercados.json
    python3 cargar_oasis.py --mercado MLC             # carga
    python3 cargar_oasis.py --todos                   # carga todos

Opciones:
    --mercado MLC     mercado a cargar; el ID del Sheets sale de mercados.json
    --todos           todos los de mercados.json, uno por uno (cada uno en su
                      transacción: si uno falla, los demás siguen)
    --sheet ID        otro Sheets para --mercado (no el de mercados.json)
    archivo.xlsx      un Excel descargado, con --mercado
    --credenciales    llave de la cuenta de servicio (default: credenciales_google.json)
    --dry-run         hace todo dentro de una transacción y la revierte
    --ensayo-sql F    aplica el SQL de F dentro de esa misma transacción antes
                      de cargar (implica --dry-run). Sirve para ver cómo quedaría
                      la carga con una migración que todavía no se aplicó.

El mercado se valida contra la celda PAÍS del Flow (C3) y contra la tabla
mercados: si no coinciden, se detiene sin tocar nada.
"""

import argparse
import json
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, timedelta

import openpyxl
import psycopg2
import psycopg2.extras

RAIZ = os.path.dirname(os.path.abspath(__file__))


# ---------------------------------------------------------------------------
# Localización de encabezados.
#
# El script NO usa números de fila ni de columna fijos. Busca la fila de
# encabezados por un texto ancla y mapea cada columna por su nombre. Si el
# Sheets cambia de layout, se adapta; si falta un encabezado requerido, se
# detiene y dice cuál.
# ---------------------------------------------------------------------------

def normalizar(v):
    """'  Sub Campaña\n(x) ' -> 'SUB CAMPANA (X)'. Sin acentos, sin saltos."""
    if v is None:
        return ""
    s = unicodedata.normalize("NFD", str(v))
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return " ".join(s.split()).upper()


def localizar(ws, ancla_col, ancla_texto, columnas, hoja):
    """
    Devuelve (fila_datos, {clave: indice_columna}).

    ancla_col    : columna (1-based) donde debe aparecer ancla_texto
    ancla_texto  : texto normalizado que identifica la fila de encabezados
    columnas     : {clave: encabezado_esperado_normalizado}
    """
    fila_hdr = None
    for r in range(1, min(ws.max_row, 50) + 1):
        if normalizar(ws.cell(row=r, column=ancla_col).value) == ancla_texto:
            fila_hdr = r
            break
    if fila_hdr is None:
        sys.exit(f"[{hoja}] No encontré la fila de encabezados: "
                 f"esperaba '{ancla_texto}' en la columna {ancla_col}.")

    encabezados = {}
    for c in range(1, ws.max_column + 1):
        h = normalizar(ws.cell(row=fila_hdr, column=c).value)
        if h and h not in encabezados:
            encabezados[h] = c

    mapa, faltan = {}, []
    for clave, esperado in columnas.items():
        idx = encabezados.get(esperado)
        if idx is None:
            # tolera encabezados con texto adicional, ej. "MATT CAMPAIGN TYPE (NO ESCRIBIR A MANO)"
            idx = next((i for h, i in encabezados.items() if h.startswith(esperado)), None)
        if idx is None:
            faltan.append(f"{clave} ('{esperado}')")
        else:
            mapa[clave] = idx

    if faltan:
        sys.exit(f"[{hoja}] Faltan encabezados en la fila {fila_hdr}: " + ", ".join(faltan)
                 + f"\n  Encabezados encontrados: {sorted(encabezados)}")

    return fila_hdr + 1, mapa


def pais_del_flow(ws):
    """El mercado que dice el Flow: la celda a la derecha de 'PAÍS' (C3 hoy)."""
    for r in range(1, 11):
        for c in range(1, 4):
            if normalizar(ws.cell(row=r, column=c).value) == "PAIS":
                for c2 in range(c + 1, c + 6):
                    v = texto(ws.cell(row=r, column=c2).value)
                    if v:
                        return normalizar(v)
                return None
    return None


FLOW_COLS = {
    "codigo": "ID", "cliente": "CLIENTE", "campana": "CAMPANA",
    "tipo_compra": "TIPO DE COMPRA", "proveedor": "PROVEEDOR", "medio": "MEDIO",
    "tipo_costo": "TIPO DE COSTO", "formato": "FORMATO", "ubicacion": "UBICACION",
    "trafico": "TRAFICO", "ciudad": "CIUDAD", "tiempo": "TIEMPO",
    "tarifa_bruta": "TARIFA BRUTA", "descuento_pct": "DESC %", "tarifa_neta": "TARIFA NETA",
    "cantidad": "CANTIDAD ELEMENTOS", "nro_semanas": "NRO. SEMANAS",
    "valor_total": "VALOR TOTAL", "fecha_inicio": "FECHA INICIO", "fecha_fin": "FECHA FIN",
}

TRACKING_COLS = {
    "compra": "ID FLOW", "material": "ID EJECUCION", "sub_campana": "SUB CAMPANA",
    "referencia": "REFERENCIA", "enlace": "ENLACE MATERIAL",
    "reporte": "REPORTE DE IMPLEMENTACION",
    "fecha_inicio": "FECHA INICIO", "fecha_fin": "FECHA FIN",
}

GLOSARIO_COLS = {
    "nombre_calendario": "NOMBRE DE CAMPANA EN CALENDARIO", "mercado": "MERCADO",
    "categorizacion": "CATEGORIZACION", "nombre_unico": "NOMBRE UNICO DE CAMPANA",
    "matt_campaign_type": "MATT CAMPAIGN TYPE",
}


# ---------------------------------------------------------------------------
# Lectura directa de Google Sheets.
#
# Presenta las hojas con la misma interfaz que openpyxl (.cell(row, col).value,
# .max_row, .max_column) para que el resto del script no distinga el origen.
# ---------------------------------------------------------------------------

_EPOCA_SHEETS = date(1899, 12, 30)   # el serial 1 de Sheets es el 31/12/1899


class _Celda:
    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


class _HojaSheets:
    def __init__(self, filas):
        self._filas = filas
        self.max_row = len(filas)
        self.max_column = max((len(f) for f in filas), default=0)

    def cell(self, row, column):
        try:
            v = self._filas[row - 1][column - 1]
        except IndexError:
            return _Celda(None)
        return _Celda(None if v == "" else v)


class _LibroSheets:
    def __init__(self, sheet_id, credenciales):
        try:
            import gspread
        except ImportError:
            sys.exit("Falta la librería gspread. Instálala con: pip3 install gspread")
        if not os.path.exists(credenciales):
            sys.exit(f"No encuentro {credenciales}. Es la llave JSON de la cuenta de servicio.")
        cliente = gspread.service_account(filename=credenciales)
        try:
            self._libro = cliente.open_by_key(sheet_id)
        except gspread.exceptions.APIError as e:
            sys.exit(f"No pude abrir el Sheets. ¿Lo compartiste con la cuenta de servicio?\n{e}")
        self._cache = {}

    def __getitem__(self, nombre):
        if nombre not in self._cache:
            try:
                hoja = self._libro.worksheet(nombre)
            except Exception:
                sys.exit(f"El Sheets no tiene una hoja llamada '{nombre}'.")
            filas = hoja.get_all_values(value_render_option="UNFORMATTED_VALUE")
            self._cache[nombre] = _HojaSheets(filas)
        return self._cache[nombre]


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------

def texto(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def espacios(v):
    """Texto con espacios limpios (sin dobles, sin bordes), o None."""
    t = texto(v)
    return " ".join(t.split()) if t else None


def normalizar_tiempo(v):
    """TIEMPO a su forma única, por la unidad que nombra. Devuelve (valor, cambió).

    Mes, Meses, MES, Mensual, '2 MESES', '1 MES Y MEDIO' → MES
    Semana, Semanas, Semanal                             → SEMANA
    Día, Dia, Dias, Diario, '21 días', '1 DIA'           → DIAS
    El valor total multiplica por semanas solo con tiempo = 'SEMANA', así que la
    unidad es lo único que importa; el número de días se pierde a propósito.
    """
    t = texto(v)
    if t is None:
        return None, False
    k = normalizar(t)
    if re.search(r"SEMAN", k):
        nuevo = "SEMANA"
    elif re.search(r"\bMES|MENSUAL", k):
        nuevo = "MES"
    elif re.search(r"\bDIA|DIARIO", k):
        nuevo = "DIAS"
    else:
        nuevo = espacios(t)
    return nuevo, nuevo != t


def numero(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^0-9.\-]", "", str(v))
    try:
        return float(s) if s not in ("", "-", ".") else None
    except ValueError:
        return None


def fecha(v):
    """Acepta datetime (Excel) o serial numérico (Sheets). Texto -> None."""
    if v is None or v == "":
        return None
    if hasattr(v, "date"):
        return v.date()
    if isinstance(v, (int, float)) and 20000 <= v <= 80000:   # ~1954 a ~2119
        return _EPOCA_SHEETS + timedelta(days=int(v))
    return None


def secuencia_desde_codigo(codigo):
    """'001-C' -> ('001', 3). Devuelve (None, None) si el formato no calza."""
    m = re.match(r"^(.+)-([A-Z])$", str(codigo).strip())
    if not m:
        return None, None
    return m.group(1), ord(m.group(2)) - 64


def filas(ws, desde, col_clave):
    for r in range(desde, ws.max_row + 1):
        if texto(ws.cell(row=r, column=col_clave).value):
            yield r


# ---------------------------------------------------------------------------
# Proveedores: el detector de duplicados de la carga, con reglas fijas.
#
# (a) Diferencias solo de mayúsculas, tildes, espacios, puntuación o sufijo
#     societario (S.A., SAS, SRL…) → misma clave → se unifica sola al nombre
#     que ya existe en la base (o al más frecuente en el Sheets si es nuevo).
# (b) Palabras de más o de menos (MASSIVA / MASSIVA METRO) o una posible
#     abreviación (JCDX / JCDECAUX, AM ADS / AMERICAN ADS) → se cargan como
#     proveedores aparte y van a duplicados_propuestos para confirmar en
#     Alertas. Nada de (b) se unifica sin clic.
# ---------------------------------------------------------------------------

SUFIJOS_SOCIETARIOS = {"SA", "SAS", "SRL", "LTDA", "LIMITADA", "SPA", "SL", "INC", "LLC",
                       "SAPI", "CV", "SACV", "SADECV", "SC", "EIRL", "SAC", "SAU"}


def tokens_proveedor(nombre):
    """'OSP S. A.' -> ['OSP']; 'TOTAL OUTDOOR S.R.L.' -> ['TOTAL', 'OUTDOOR']; 'J&M' -> ['JM']."""
    s = normalizar(nombre).replace(".", "")
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    # Letras sueltas seguidas se juntan: 'S A' -> 'SA', 'S A S' -> 'SAS', 'J M' -> 'JM'.
    tok, de_sueltas = [], False
    for t in s.split():
        if len(t) == 1 and de_sueltas:
            tok[-1] += t
        else:
            tok.append(t)
            de_sueltas = len(t) == 1
    while len(tok) > 1 and (tok[-1] in SUFIJOS_SOCIETARIOS or tok[-1] == "DE"):
        tok.pop()
    return tok


def clave_proveedor(nombre):
    return "".join(tokens_proveedor(nombre)) or re.sub(r"[^A-Z0-9]", "", normalizar(nombre))


def _subsecuencia(corta, larga):
    it = iter(larga)
    return all(ch in it for ch in corta)


def parecido_b(a, b):
    """Motivo si a y b parecen el mismo proveedor por la regla (b), o None."""
    ta, tb = tokens_proveedor(a), tokens_proveedor(b)
    sa, sb = set(ta), set(tb)
    if sa and sb and sa != sb and (sa < sb or sb < sa):
        return "palabras de más o de menos"
    if len(ta) == len(tb) and ta != tb:
        pares = list(zip(ta, tb))
        def abrevia(x, y):
            if x == y:
                return True
            corta, larga = sorted((x, y), key=len)
            return len(corta) >= 2 and corta[0] == larga[0] and len(corta) < len(larga) and _subsecuencia(corta, larga)
        if all(abrevia(x, y) for x, y in pares):
            return "posible abreviación"
    return None


# ---------------------------------------------------------------------------
# Carga de un mercado
# ---------------------------------------------------------------------------

class Carga:
    def __init__(self, conn, wb, mercado_codigo, dry_run=False):
        self.cur = conn.cursor()
        self.conn = conn
        self.wb = wb
        self.mercado_codigo = mercado_codigo
        self.dry_run = dry_run
        self.avisos = defaultdict(list)      # se listan con los primeros 10
        self.notas = Counter()               # normalizaciones en silencio: solo conteo
        self.stats = defaultdict(int)

        self.flow_ini, self.FLOW = localizar(wb["Flow"], 1, "ID", FLOW_COLS, "Flow")
        self.trk_ini,  self.TRK  = localizar(wb["Tracking"], 1, "ID FLOW", TRACKING_COLS, "Tracking")
        self.glo_ini,  self.GLO  = localizar(wb["Glosario"], 3, "MERCADO", GLOSARIO_COLS, "Glosario")

        self.cur.execute("SELECT id, moneda FROM mercados WHERE codigo = %s", (mercado_codigo,))
        fila = self.cur.fetchone()
        if not fila:
            sys.exit(f"El mercado {mercado_codigo} no existe en la tabla mercados.")
        self.mercado_id, self.moneda = fila

        # Valores que admite cada enum, leídos de la base.
        self.enums = {}
        for col, tipo in (("tipo_compra", "tipo_compra_enum"), ("medio", "medio_enum"), ("tipo_costo", "tipo_costo_enum")):
            self.cur.execute("SELECT enumlabel FROM pg_enum WHERE enumtypid = %s::regtype", (tipo,))
            self.enums[col] = {r[0] for r in self.cur.fetchall()}

        self.cur.execute("SELECT upper(btrim(texto_flow)), taxonomia FROM campanas_alias WHERE mercado_id = %s",
                         (self.mercado_id,))
        self.alias = dict(self.cur.fetchall())
        self.unificaciones()

    # -- unificaciones -----------------------------------------------------

    def unificaciones(self):
        """Duplicados que alguien unificó en Alertas (sql/duplicados.sql) o que
        la carga unificó sola (regla a).

        El Sheets sigue trayendo los textos viejos: sin esto, cada carga desharía
        la unificación. Proveedor: nombre → nombre canónico. Soporte: (proveedor
        ya traducido, ubicación, ciudad) → ubicación canónica.
        """
        self.cur.execute("""
            SELECT tipo, valores, canonico FROM duplicados_propuestos
            WHERE mercado_id = %s AND estado = 'UNIFICADO'
            ORDER BY resuelto_en, id
        """, (self.mercado_id,))
        unificados = self.cur.fetchall()

        self.unif_proveedor = {}
        for tipo, valores, canonico in unificados:
            if tipo == "PROVEEDOR":
                for v in valores:
                    if v["texto"] != canonico:
                        self.unif_proveedor[v["texto"]] = canonico

        self.unif_soporte = {}
        for tipo, valores, canonico in unificados:
            if tipo == "SOPORTE":
                for v in valores:
                    prov = self.proveedor_unificado(v["proveedor"])
                    if v["ubicacion"] != canonico:
                        self.unif_soporte[(prov, v["ubicacion"], v["ciudad"])] = canonico

    # Siguen la cadena: si A se unificó en B y luego B en C, A termina en C.
    def proveedor_unificado(self, nombre):
        vistos = set()
        while nombre in self.unif_proveedor and nombre not in vistos:
            vistos.add(nombre)
            nombre = self.unif_proveedor[nombre]
        return nombre

    def ubicacion_unificada(self, proveedor, ubicacion, ciudad):
        vistos = set()
        while (proveedor, ubicacion, ciudad) in self.unif_soporte and ubicacion not in vistos:
            vistos.add(ubicacion)
            ubicacion = self.unif_soporte[(proveedor, ubicacion, ciudad)]
        return ubicacion

    # -- filas del Flow ----------------------------------------------------

    def filas_flow(self):
        """Filas del Flow con ID y algo más. Las que solo tienen el ID (la
        plantilla numera filas vacías) se saltan y se cuentan."""
        ws = self.wb["Flow"]
        g = lambda r, k: ws.cell(row=r, column=self.FLOW[k]).value
        utiles = []
        for r in filas(ws, self.flow_ini, self.FLOW["codigo"]):
            if all(texto(g(r, k)) is None for k in ("proveedor", "formato", "ubicacion", "tarifa_neta",
                                                     "valor_total", "fecha_inicio", "campana")):
                self.stats["filas_vacias"] += 1
                continue
            utiles.append(r)
        return utiles

    # -- campañas ----------------------------------------------------------

    def campanas(self):
        ws = self.wb["Glosario"]
        datos = []
        for r in filas(ws, self.glo_ini, self.GLO["nombre_unico"]):
            mercado = texto(ws.cell(row=r, column=self.GLO["mercado"]).value)
            nombre = texto(ws.cell(row=r, column=self.GLO["nombre_unico"]).value)
            if not mercado or not nombre:
                continue
            datos.append((
                mercado.upper(), nombre,
                texto(ws.cell(row=r, column=self.GLO["nombre_calendario"]).value),
                texto(ws.cell(row=r, column=self.GLO["categorizacion"]).value),
                texto(ws.cell(row=r, column=self.GLO["matt_campaign_type"]).value),
            ))

        vistos = set()
        unicos = []
        for d in datos:
            if (d[0], d[1]) in vistos:
                continue
            vistos.add((d[0], d[1]))
            unicos.append(d)

        psycopg2.extras.execute_batch(self.cur, """
            INSERT INTO campanas
                (mercado_id, nombre_unico, nombre_calendario,
                 categorizacion, matt_campaign_type, sincronizado_en)
            SELECT m.id, %s, %s, %s, %s, now()
            FROM mercados m WHERE m.codigo = %s
            ON CONFLICT (mercado_id, nombre_unico) DO UPDATE SET
                nombre_calendario  = EXCLUDED.nombre_calendario,
                categorizacion     = EXCLUDED.categorizacion,
                matt_campaign_type = EXCLUDED.matt_campaign_type,
                sincronizado_en    = now()
        """, [(d[1], d[2], d[3], d[4], d[0]) for d in unicos])

        self.stats["campanas"] = len(unicos)

        self.cur.execute("SELECT c.nombre_unico, c.id FROM campanas c WHERE c.mercado_id = %s", (self.mercado_id,))
        self.mapa_campanas = dict(self.cur.fetchall())

    def resolver_campanas(self, textos):
        """{texto del Flow: (taxonomía o None, campana_id o None, de dónde salió)}.

        Las reglas viven en la base (fn_taxonomia_de_texto, fn_campana_de_taxonomia)
        para que la carga y la corrección en Alertas den lo mismo.
        """
        textos = sorted({t for t in textos if t})
        self.cur.execute("""
            SELECT t, public.fn_taxonomia_de_texto(%s, t) FROM unnest(%s::text[]) t
        """, (self.mercado_codigo, textos))
        tax = dict(self.cur.fetchall())
        res = {}
        for t in textos:
            if tax.get(t):
                origen = "recortada" if tax[t] != t.strip().upper() else "flow"
                res[t] = [tax[t], origen]
            elif t.strip().upper() in self.alias:
                res[t] = [self.alias[t.strip().upper()], "alias"]
            else:
                res[t] = [None, None]
        validas = sorted({v[0] for v in res.values() if v[0]})
        self.cur.execute("""
            SELECT x, public.fn_campana_de_taxonomia(%s::smallint, x) FROM unnest(%s::text[]) x
        """, (self.mercado_id, validas))
        ids = dict(self.cur.fetchall())
        return {t: (v[0], ids.get(v[0]) if v[0] else None, v[1]) for t, v in res.items()}

    # -- proveedores -------------------------------------------------------

    def proveedores(self, filas_flow):
        """Catálogo de proveedores del mercado, derivado del Flow, con la regla (a)."""
        ws = self.wb["Flow"]
        frecuencia = Counter()
        for r in filas_flow:
            n = espacios(ws.cell(row=r, column=self.FLOW["proveedor"]).value)
            if n:
                frecuencia[self.proveedor_unificado(n.upper())] += 1

        # Proveedores que ya están en la base, por clave: gana el activo con más compras.
        self.cur.execute("""
            SELECT p.id, p.nombre, p.activo, count(c.id)
            FROM proveedores p LEFT JOIN compras c ON c.proveedor_id = p.id
            WHERE p.mercado_id = %s GROUP BY p.id ORDER BY p.activo DESC, count(c.id) DESC, p.id
        """, (self.mercado_id,))
        existentes = {}
        for pid, nombre, activo, n in self.cur.fetchall():
            existentes.setdefault(clave_proveedor(nombre), nombre)

        grupos = defaultdict(list)
        for nombre in frecuencia:
            grupos[clave_proveedor(nombre)].append(nombre)

        self.canonico = {}          # nombre del Sheets (ya en mayúsculas) → nombre canónico
        self.unificados_a = []      # (canónico, [variantes]) de la regla (a)
        for k, nombres in grupos.items():
            canon = existentes.get(k) or max(nombres, key=lambda n: (frecuencia[n], len(n), n))
            for n in nombres:
                self.canonico[n] = canon
            variantes = sorted(n for n in nombres if n != canon)
            if variantes:
                self.unificados_a.append((canon, variantes))
                self.notas["proveedores unificados por la regla (a)"] += len(variantes)

        nuevos = sorted(set(self.canonico.values()))
        psycopg2.extras.execute_batch(self.cur, """
            INSERT INTO proveedores (mercado_id, nombre) VALUES (%s, %s)
            ON CONFLICT (mercado_id, nombre) DO UPDATE SET activo = true
        """, [(self.mercado_id, n) for n in nuevos])
        self.stats["proveedores"] = len(nuevos)

        self.cur.execute("SELECT nombre, id FROM proveedores WHERE mercado_id = %s", (self.mercado_id,))
        self.mapa_proveedores = dict(self.cur.fetchall())

    def cerrar_proveedores(self):
        """Después de cargar las compras: deja rastro de la regla (a), apaga los
        proveedores absorbidos que quedaron sin compras y propone la regla (b)."""
        self.cur.execute("""
            SELECT valores, estado FROM duplicados_propuestos WHERE mercado_id = %s AND tipo = 'PROVEEDOR'
        """, (self.mercado_id,))
        previas = [({v["texto"] for v in valores}, estado) for valores, estado in self.cur.fetchall()]

        def compras_de(nombre):
            pid = self.mapa_proveedores.get(nombre)
            if not pid:
                return 0
            self.cur.execute("SELECT count(*) FROM compras WHERE proveedor_id = %s", (pid,))
            return self.cur.fetchone()[0]

        # (a): una propuesta UNIFICADO por grupo (así las cargas siguientes y la
        # interfaz la ven) y una fila de auditoría por variante. Solo la primera vez.
        for canon, variantes in self.unificados_a:
            textos = {canon, *variantes}
            if any(textos <= prev for prev, estado in previas if estado == "UNIFICADO"):
                continue
            valores = [{"texto": t, "proveedor_id": self.mapa_proveedores.get(t), "compras": compras_de(t)}
                       for t in sorted(textos)]
            self.cur.execute("""
                INSERT INTO duplicados_propuestos (mercado_id, tipo, valores, canonico, confianza, razon, estado, resuelto_en)
                VALUES (%s, 'PROVEEDOR', %s, %s, 'ALTA', %s, 'UNIFICADO', now())
            """, (self.mercado_id, json.dumps(valores, ensure_ascii=False), canon,
                  "Unificado en la carga: solo cambian mayúsculas, tildes, espacios, puntuación o sufijo societario."))
            for v in variantes:
                self.cur.execute("""
                    INSERT INTO auditoria (tabla, fila_id, campo, valor_anterior, valor_nuevo)
                    VALUES ('proveedores', %s, '(unificado en carga)', %s, %s)
                """, (self.mapa_proveedores[canon], v, canon))
            previas.append((textos, "UNIFICADO"))
            self.avisos["proveedores unificados en la carga (regla a)"].append(f"{' / '.join(variantes)} → {canon}")

        # Propuestas PENDIENTE de antes que la regla (a) ya resolvió: se cierran.
        self.cur.execute("""
            SELECT id, valores FROM duplicados_propuestos
            WHERE mercado_id = %s AND tipo = 'PROVEEDOR' AND estado = 'PENDIENTE'
        """, (self.mercado_id,))
        for pid, valores in self.cur.fetchall():
            claves = {clave_proveedor(v["texto"]) for v in valores}
            if len(claves) == 1:
                canon = self.canonico.get(valores[0]["texto"]) or valores[0]["texto"]
                self.cur.execute("""
                    UPDATE duplicados_propuestos SET estado = 'UNIFICADO', canonico = %s, resuelto_en = now(),
                        razon = coalesce(razon, '') || ' (resuelta en la carga: misma clave)'
                    WHERE id = %s
                """, (canon, pid))

        # Absorbidos sin compras: inactivos, como al unificar en Alertas. Son los
        # de la regla (a) y los de una unificación UNIFICADO (de Alertas o por
        # decisión), que la carga acaba de dejar sin compras.
        absorbidos = {v for _, vs in self.unificados_a for v in vs} | set(self.unif_proveedor)
        self.cur.execute("""
            UPDATE proveedores p SET activo = false
            WHERE p.mercado_id = %s AND p.activo AND p.nombre = ANY (%s)
              AND NOT EXISTS (SELECT 1 FROM compras c WHERE c.proveedor_id = p.id)
            RETURNING p.id, p.nombre
        """, (self.mercado_id, sorted(absorbidos)))
        apagados = self.cur.fetchall()
        if apagados:
            self.cur.execute("""
                SELECT u.correo, p.nombre FROM usuarios u JOIN proveedores p ON p.id = u.proveedor_id
                WHERE u.proveedor_id = ANY (%s) AND u.activo
            """, ([a[0] for a in apagados],))
            for correo, nombre in self.cur.fetchall():
                self.avisos["usuario PROVEEDOR apunta a un proveedor absorbido (cambiar su proveedor)"].append(
                    f"{correo}: {nombre} → {self.canonico.get(nombre, '?')}")

        # (b): grupos por pares parecidos, sobre los canónicos que tienen compras.
        nombres = sorted(n for n in set(self.canonico.values()) if compras_de(n))
        padre = {n: n for n in nombres}
        def raiz(x):
            while padre[x] != x:
                padre[x] = padre[padre[x]]
                x = padre[x]
            return x
        motivos = {}
        for i, a in enumerate(nombres):
            for b in nombres[i + 1:]:
                m = parecido_b(a, b)
                if m:
                    padre[raiz(a)] = raiz(b)
                    motivos.setdefault(raiz(b), set()).add(m)
        grupos = defaultdict(list)
        for n in nombres:
            grupos[raiz(n)].append(n)
        for r, g in grupos.items():
            if len(g) < 2:
                continue
            textos = set(g)
            # No repetir lo ya decidido ni encimar una propuesta pendiente.
            if any(len(textos & prev) >= 2 for prev, _ in previas):
                continue
            valores = sorted(({"texto": t, "proveedor_id": self.mapa_proveedores[t], "compras": compras_de(t)} for t in g),
                             key=lambda v: (-v["compras"], v["texto"]))
            razon = "Detectado en la carga: " + " y ".join(sorted(set().union(*[motivos.get(raiz(t), set()) for t in g]) or {"parecidos"})) \
                    + ". Confirma si son el mismo proveedor."
            self.cur.execute("""
                INSERT INTO duplicados_propuestos (mercado_id, tipo, valores, canonico, confianza, razon, estado)
                VALUES (%s, 'PROVEEDOR', %s, %s, 'MEDIA', %s, 'PENDIENTE')
            """, (self.mercado_id, json.dumps(valores, ensure_ascii=False), valores[0]["texto"], razon))
            previas.append((textos, "PENDIENTE"))

    # -- compras --------------------------------------------------------

    def enum_o_fuera(self, col, valor, fuera):
        """Valor del enum, o None (y lo original en fuera_de_catalogo) si no está."""
        t = texto(valor)
        if t is None:
            return None
        k = normalizar(t)
        if k in self.enums[col]:
            return k
        fuera[col] = t
        return None

    def compras(self, filas_flow):
        ws = self.wb["Flow"]
        g = lambda r, k: ws.cell(row=r, column=self.FLOW[k]).value
        camp = self.resolver_campanas([texto(g(r, "campana")) for r in filas_flow])
        pendientes = []

        for r in filas_flow:
            codigo = texto(g(r, "codigo"))
            fuera = {}

            # Fechas: las dos o ninguna; invertidas entran sin fechas (alerta).
            fi, ff = fecha(g(r, "fecha_inicio")), fecha(g(r, "fecha_fin"))
            if (fi is None) != (ff is None):
                self.avisos["compra con una sola fecha (entra sin fechas)"].append(codigo)
                fi = ff = None
            elif fi and ff < fi:
                self.avisos["compra con fechas invertidas (entra sin fechas)"].append(f"{codigo}: {fi} → {ff}")
                fi = ff = None

            # Campaña: taxonomía del Flow, o su equivalencia en campanas_alias.
            campana_texto = texto(g(r, "campana"))
            taxonomia, campana_id, origen = camp.get(campana_texto, (None, None, None))
            if origen == "recortada":
                self.notas["campaña con texto después de la taxonomía (se tomó la taxonomía)"] += 1
            elif origen == "alias":
                self.notas["campaña traducida por campanas_alias"] += 1

            tipo_costo = self.enum_o_fuera("tipo_costo", g(r, "tipo_costo"), fuera)
            tipo_compra = self.enum_o_fuera("tipo_compra", g(r, "tipo_compra"), fuera)
            medio = self.enum_o_fuera("medio", g(r, "medio"), fuera)

            descuento = numero(g(r, "descuento_pct")) or 0
            if not 0 <= descuento <= 1:
                self.avisos["compra con descuento fuera de 0–1 (entra con 0)"].append(f"{codigo}: {descuento}")
                descuento = 0

            tiempo, cambio = normalizar_tiempo(g(r, "tiempo"))
            if cambio:
                self.notas["tiempo normalizado"] += 1

            formato = espacios(g(r, "formato"))
            formato = formato.upper() if formato else None
            if formato and formato != texto(g(r, "formato")):
                self.notas["formato a mayúsculas o espacios limpios"] += 1

            # valor_total_manual: el valor no sale de la fórmula. Con TIEMPO = SEMANA
            # (y tipo de costo distinto de PRODUCCION) la tarifa es semanal:
            # tarifa neta × cantidad × nro. semanas (vacío = 0); si no, tarifa
            # neta × cantidad. Sin tipo de costo no hay fórmula: vale lo que diga
            # el Sheets (0 si no dice nada) y se marca como manual.
            tarifa_neta = numero(g(r, "tarifa_neta")) or 0
            cantidad = int(numero(g(r, "cantidad")) or 1)
            valor_total = numero(g(r, "valor_total")) or 0
            if tipo_costo is None:
                valor_total_manual = True
            else:
                por_semana = tiempo == "SEMANA" and tipo_costo != "PRODUCCION"
                semanas = (numero(g(r, "nro_semanas")) or 0) if por_semana else 1
                valor_total_manual = abs(valor_total - round(tarifa_neta * cantidad * semanas, 2)) > 1

            prov = None
            prov_sheets = espacios(g(r, "proveedor"))
            if prov_sheets:
                traducido = self.proveedor_unificado(prov_sheets.upper())
                prov = self.canonico.get(traducido, traducido)
                if traducido != prov_sheets.upper():
                    self.notas["proveedor traducido por una unificación de Alertas"] += 1
            proveedor_id = self.mapa_proveedores.get(prov) if prov else None

            ubicacion_sheets = espacios(g(r, "ubicacion"))
            ciudad = espacios(g(r, "ciudad"))
            ubicacion = self.ubicacion_unificada(prov, ubicacion_sheets, ciudad) if ubicacion_sheets else None
            if ubicacion != ubicacion_sheets:
                self.notas["ubicación traducida por una unificación de Alertas"] += 1

            pendientes.append((
                self.mercado_id, codigo, espacios(g(r, "cliente")),
                campana_id, campana_texto, taxonomia, proveedor_id,
                tipo_compra, medio, tipo_costo,
                formato, ubicacion, ciudad,
                numero(g(r, "trafico")), tiempo,
                numero(g(r, "tarifa_bruta")) or 0, descuento,
                tarifa_neta, cantidad,
                numero(g(r, "nro_semanas")), valor_total, valor_total_manual,
                fi, ff, json.dumps(fuera, ensure_ascii=False) if fuera else None,
            ))

        psycopg2.extras.execute_batch(self.cur, """
            INSERT INTO compras (
                mercado_id, codigo, cliente, campana_id, campana_texto, campana_taxonomia, proveedor_id,
                tipo_compra, medio, tipo_costo, formato, ubicacion, ciudad,
                trafico, tiempo, tarifa_bruta, descuento_pct, tarifa_neta,
                cantidad, nro_semanas, valor_total, valor_total_manual,
                fecha_inicio, fecha_fin, fuera_de_catalogo)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s::tipo_compra_enum,%s::medio_enum,%s::tipo_costo_enum,
                    %s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s::jsonb)
            ON CONFLICT (mercado_id, codigo) DO UPDATE SET
                cliente           = EXCLUDED.cliente,
                campana_id        = EXCLUDED.campana_id,
                campana_texto     = EXCLUDED.campana_texto,
                campana_taxonomia = EXCLUDED.campana_taxonomia,
                proveedor_id      = EXCLUDED.proveedor_id,
                tipo_compra       = EXCLUDED.tipo_compra,
                medio             = EXCLUDED.medio,
                tipo_costo        = EXCLUDED.tipo_costo,
                formato           = EXCLUDED.formato,
                ubicacion         = EXCLUDED.ubicacion,
                ciudad            = EXCLUDED.ciudad,
                trafico           = EXCLUDED.trafico,
                tiempo            = EXCLUDED.tiempo,
                tarifa_bruta      = EXCLUDED.tarifa_bruta,
                descuento_pct     = EXCLUDED.descuento_pct,
                tarifa_neta       = EXCLUDED.tarifa_neta,
                cantidad          = EXCLUDED.cantidad,
                nro_semanas       = EXCLUDED.nro_semanas,
                valor_total       = EXCLUDED.valor_total,
                valor_total_manual = EXCLUDED.valor_total_manual,
                fecha_inicio      = EXCLUDED.fecha_inicio,
                fecha_fin         = EXCLUDED.fecha_fin,
                fuera_de_catalogo = EXCLUDED.fuera_de_catalogo,
                actualizado_en    = now()
        """, pendientes)

        self.stats["compras_leidas"] = len(filas_flow) + self.stats["filas_vacias"]
        self.stats["compras"] = len(pendientes)

        self.cur.execute("SELECT codigo, id FROM compras WHERE mercado_id = %s", (self.mercado_id,))
        self.mapa_compras = dict(self.cur.fetchall())

    # -- materiales -------------------------------------------------------

    def materiales(self):
        ws = self.wb["Tracking"]
        pendientes = []
        leidos = 0

        for r in filas(ws, self.trk_ini, self.TRK["material"]):
            g = lambda k: ws.cell(row=r, column=self.TRK[k]).value
            codigo = texto(g("material"))
            leidos += 1

            mat_codigo, secuencia = secuencia_desde_codigo(codigo)
            if secuencia is None:
                self.avisos["material con código de formato inesperado (no entra)"].append(codigo)
                continue

            compra_id = self.mapa_compras.get(mat_codigo)
            if not compra_id:
                self.avisos["material cuya compra no está en el Flow (no entra)"].append(codigo)
                continue

            # La sub campaña puede traer valores fuera del glosario (ODM,
            # espacios al final). Entran en nulo y quedan en v_alertas.
            sub = texto(g("sub_campana"))
            sub_id = self.mapa_campanas.get(sub) if sub else None
            if sub and not sub_id:
                self.avisos["sub campaña fuera del glosario (entra sin sub campaña)"].append(f"{codigo}: {sub!r}")

            raw_fi, raw_ff = g("fecha_inicio"), g("fecha_fin")
            fi, ff = fecha(raw_fi), fecha(raw_ff)
            if (fi is None) != (ff is None):
                self.avisos["material con una sola fecha (entra sin fechas)"].append(codigo)
                fi = ff = None
            elif fi is None and (texto(raw_fi) or texto(raw_ff)):
                self.avisos["material con fecha escrita como texto (entra sin fechas)"].append(
                    f"{codigo}: {texto(raw_fi)!r} → {texto(raw_ff)!r}")
            if fi and ff and ff < fi:
                self.avisos["material con fechas invertidas (entra sin fechas)"].append(f"{codigo}: {fi} → {ff}")
                fi = ff = None

            pendientes.append((
                compra_id, secuencia, sub_id,
                texto(g("referencia")), texto(g("enlace")), texto(g("reporte")),
                fi, ff,
            ))

        psycopg2.extras.execute_batch(self.cur, """
            INSERT INTO materiales (
                compra_id, secuencia, sub_campana_id,
                referencia, enlace, reporte_implementacion,
                fecha_inicio, fecha_fin)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s)
            ON CONFLICT (compra_id, secuencia) DO UPDATE SET
                sub_campana_id         = EXCLUDED.sub_campana_id,
                referencia             = EXCLUDED.referencia,
                enlace                 = EXCLUDED.enlace,
                reporte_implementacion = EXCLUDED.reporte_implementacion,
                fecha_inicio           = EXCLUDED.fecha_inicio,
                fecha_fin              = EXCLUDED.fecha_fin,
                actualizado_en         = now()
        """, pendientes)

        self.stats["materiales_leidos"] = leidos
        self.stats["materiales"] = len(pendientes)

    # -- huérfanos ---------------------------------------------------------

    def huerfanos(self):
        """Filas que existen en la base pero ya no en el Sheets. No se borran solas."""
        ws = self.wb["Flow"]
        codigos = {texto(ws.cell(row=r, column=self.FLOW["codigo"]).value)
                   for r in filas(ws, self.flow_ini, self.FLOW["codigo"])}
        self.cur.execute("SELECT codigo FROM compras WHERE mercado_id = %s", (self.mercado_id,))
        for (c,) in self.cur.fetchall():
            if c not in codigos:
                self.avisos["compra en la base que ya no está en el Sheets (no se borra)"].append(c)

        wt = self.wb["Tracking"]
        en_sheets = set()
        for r in filas(wt, self.trk_ini, self.TRK["material"]):
            mat, sec = secuencia_desde_codigo(texto(wt.cell(row=r, column=self.TRK["material"]).value))
            if sec:
                en_sheets.add((mat, sec))
        self.cur.execute("""
            SELECT m.codigo, e.secuencia
            FROM materiales e JOIN compras m ON m.id = e.compra_id
            WHERE m.mercado_id = %s
        """, (self.mercado_id,))
        for mat, sec in self.cur.fetchall():
            if (mat, sec) not in en_sheets:
                self.avisos["material en la base que ya no está en el Sheets (no se borra)"].append(
                    f"{mat}-{chr(64 + sec)}")

    # -- alertas como las verá la pantalla --------------------------------

    def alertas(self):
        self.cur.execute("""
            SELECT alerta, count(*),
                   (array_agg(referencia || coalesce(' · ' || left(detalle, 60), '') ORDER BY referencia))[1:10]
            FROM v_alertas WHERE mercado_id = %s GROUP BY alerta ORDER BY alerta
        """, (self.mercado_id,))
        self.lista_alertas = self.cur.fetchall()

    # -- reporte -----------------------------------------------------------

    def reporte(self):
        s = self.stats
        print("\n" + "=" * 70)
        print(f"  {self.mercado_codigo} ({self.moneda})" + ("  —  SIMULACIÓN, nada quedó guardado" if self.dry_run else "  —  CARGADO"))
        print("=" * 70)
        print(f"  Compras:     {s['compras_leidas']:>5} leídas en el Flow · {s['filas_vacias']} filas solo con ID saltadas"
              f" · {s['compras']} entran")
        print(f"  Materiales:  {s['materiales_leidos']:>5} leídos en el Tracking · {s['materiales']} entran")
        print(f"  Proveedores: {s['proveedores']:>5} · glosario: {s['campanas']} campañas (todos los mercados)")
        if s["inversion"]:
            print(f"  Inversión:   {s['inversion']:>5} filas diarias recalculadas")

        if self.notas:
            print("\n  Normalizado en silencio")
            for k, n in sorted(self.notas.items()):
                print(f"    {k}: {n}")

        print("\n  Alertas en Alertas de " + self.mercado_codigo + " (v_alertas)")
        if not self.lista_alertas:
            print("    ninguna")
        for alerta, n, ejemplos in self.lista_alertas:
            print(f"\n    {alerta} ({n})")
            for e in ejemplos:
                print(f"        {e}")
            if n > 10:
                print(f"        … y {n - 10} más")

        if self.avisos:
            print("\n  Para revisar (no son alertas de la pantalla)")
            for motivo, items in sorted(self.avisos.items()):
                print(f"\n    {motivo} ({len(items)})")
                for i in items[:10]:
                    print(f"        {i}")
                if len(items) > 10:
                    print(f"        … y {len(items) - 10} más")
        print()


# ---------------------------------------------------------------------------

def cargar_mercado(mercado, wb, dry_run, ensayo_sql):
    """Carga un mercado en su propia transacción. Devuelve True si terminó bien."""
    # Todo lo del Sheets se lee antes de abrir la transacción: así los bloqueos
    # en la base duran segundos, no lo que tarde Google.
    for hoja in ("Flow", "Tracking", "Glosario"):
        wb[hoja]
    pais = pais_del_flow(wb["Flow"])
    if pais != mercado:
        print(f"  [{mercado}] El Flow dice PAÍS = {pais!r}, no {mercado}. No se cargó nada.")
        return False

    conn = psycopg2.connect()
    try:
        if ensayo_sql:
            sql = open(ensayo_sql, encoding="utf-8").read()
            sql = re.sub(r"^\s*(BEGIN|COMMIT);\s*$", "", sql, flags=re.M)
            conn.cursor().execute(sql)
        # La inversión precalculada (sql/inversion_precalculada.sql) se recalcula
        # una vez al final, no una vez por cada material que se escribe.
        cur = conn.cursor()
        cur.execute("SELECT to_regproc('public.fn_recalcular_inversion_mercado') IS NOT NULL")
        precalculada = cur.fetchone()[0]
        if precalculada:
            cur.execute("SELECT set_config('oasis.sin_recalculo', 'on', true)")
        carga = Carga(conn, wb, mercado, dry_run)
        carga.campanas()
        utiles = carga.filas_flow()
        carga.proveedores(utiles)
        carga.compras(utiles)
        carga.cerrar_proveedores()
        carga.materiales()
        carga.huerfanos()
        if precalculada:
            cur.execute("SELECT set_config('oasis.sin_recalculo', 'off', true)")
            cur.execute("SELECT public.fn_recalcular_inversion_mercado(%s::smallint)", (carga.mercado_id,))
            carga.stats["inversion"] = cur.fetchone()[0]
        carga.alertas()

        if dry_run:
            conn.rollback()
        else:
            conn.commit()
        carga.reporte()
        return True
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("archivo", nargs="?", help="Excel exportado del Sheets (con --mercado)")
    ap.add_argument("--mercado", help="mercado a cargar (MCO, MLC, …)")
    ap.add_argument("--todos", action="store_true", help="todos los de mercados.json")
    ap.add_argument("--sheet", help="ID de otro Google Sheets para --mercado")
    ap.add_argument("--credenciales", default=os.path.join(RAIZ, "credenciales_google.json"))
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--ensayo-sql", help="SQL a aplicar dentro de la transacción (implica --dry-run)")
    args = ap.parse_args()

    if bool(args.mercado) == bool(args.todos):
        ap.error("indica --mercado XXX o --todos")
    if not os.getenv("PGHOST"):
        sys.exit("Faltan las variables de conexión. Corre: source .env")
    dry_run = args.dry_run or bool(args.ensayo_sql)

    hojas = json.load(open(os.path.join(RAIZ, "mercados.json"), encoding="utf-8"))
    hojas = {k: v for k, v in hojas.items() if not k.startswith("_")}
    mercados = sorted(hojas) if args.todos else [args.mercado.upper()]

    fallas = []
    for mk in mercados:
        if args.archivo:
            wb = openpyxl.load_workbook(args.archivo, data_only=True)
        else:
            sid = args.sheet or hojas.get(mk)
            if not sid:
                sys.exit(f"{mk} no está en mercados.json; agrégalo o usa --sheet ID.")
            print(f"  [{mk}] Leyendo Google Sheets {sid} en vivo…")
            wb = _LibroSheets(sid, args.credenciales)
        try:
            if not cargar_mercado(mk, wb, dry_run, args.ensayo_sql):
                fallas.append(mk)
        except Exception as e:
            print(f"  [{mk}] FALLÓ, no se cargó nada de este mercado: {e}")
            fallas.append(mk)
            if not args.todos:
                raise

    if len(mercados) > 1:
        print("Resumen: " + ", ".join(f"{m} {'FALLÓ' if m in fallas else 'ok'}" for m in mercados))
    sys.exit(1 if fallas else 0)


if __name__ == "__main__":
    main()
